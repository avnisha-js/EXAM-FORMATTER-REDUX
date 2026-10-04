"""Math 10 certification uses the real reference and does not format."""

import shutil
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from app.reference_registry import ReferenceRegistry
from app.web import create_app, duplicate_message
from common.hashing import sha256_file
from subjects.math_10.reference.read import academic_text, read_exam, render_label, resolved_labels
from subjects.math_10.service import certify_reference, format_exam

ROOT = Path(__file__).resolve().parents[3]
REFERENCE = ROOT / "input" / "math_10" / "reference" / "Reference Exam Math Grade 10.docx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def test_reference_certifies_and_preserves_content():
    before = sha256_file(REFERENCE)
    result = certify_reference(REFERENCE)
    assert result.ok, result.message
    assert result.kind == "certified"
    assert sha256_file(REFERENCE) == before
    exam = read_exam(REFERENCE)
    labels = resolved_labels(exam)
    text = academic_text(exam)
    assert "(a)" in labels and "(b)" in labels and "(c)" in labels and "(d)" in labels
    assert "(i)" in labels and "(ii)" in labels
    assert "1." in labels
    for character in ("−", "²", "₁", "₂", "×", "√", "α", "β", "≠", "¼", "π"):
        assert character in text
    assert "x² − 49" in text
    assert "a₁/a₂ = b₁/b₂ = c₁/c₂" in text
    assert "6 × 5 = 30" in text
    assert exam.unsupported == ()


def test_list_labels_follow_the_reference_definitions():
    assert render_label("lowerLetter", "(%1)", 1) == "(a)"
    assert render_label("lowerLetter", "(%1)", 4) == "(d)"
    assert render_label("lowerRoman", "(%1)", 2) == "(ii)"
    assert render_label("decimal", "%1.", 6) == "6."
    assert render_label("bullet", "•", 1) == "bullet"


def test_omml_table_and_drawing_fail_without_being_removed(tmp_path):
    omml = tmp_path / "omml.docx"
    _inject(REFERENCE, omml, f'<m:oMath xmlns:m="{M}"><m:r><m:t>x</m:t></m:r></m:oMath>')
    drawing = tmp_path / "drawing.docx"
    _inject(REFERENCE, drawing, f'<w:r xmlns:w="{W}"><w:drawing/></w:r>')
    table = tmp_path / "table.docx"
    document = Document()
    document.add_table(rows=1, cols=1)
    document.save(str(table))
    for path, token in ((omml, "OMML"), (drawing, "drawing"), (table, "table")):
        result = certify_reference(path)
        assert result.ok is False
        assert "Reference Certification Failed." in result.message
        assert token in result.message
        assert "not removed" in result.message or "Nothing was removed." in result.message


def test_failed_reference_does_not_replace_a_certified_one(tmp_path):
    app = create_app(tmp_path)
    client = app.test_client()
    installed = client.post(
        "/reference",
        data={"grade": "10", "subject": "Math", "reference": (REFERENCE.open("rb"), "Reference Exam Math Grade 10.docx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "Reference installed successfully." in installed.get_data(as_text=True)
    assert "Grade 10 / Math" in installed.get_data(as_text=True)
    registry = ReferenceRegistry(tmp_path)
    original = sha256_file(registry.file_path(registry.find("10", "Math")))
    bad = tmp_path / "bad.docx"
    document = Document()
    document.add_paragraph("not a math reference")
    document.save(str(bad))
    rejected = client.post(
        "/reference",
        data={"grade": "10", "subject": "math", "reference": (bad.open("rb"), "bad.docx")},
        content_type="multipart/form-data",
    )
    assert duplicate_message(registry.find("10", "Math")) in rejected.get_data(as_text=True)
    current = registry.find("10", "MATH")
    assert current["filename"] == "Reference Exam Math Grade 10.docx"
    assert current["certification_status"] == "certified"
    assert current["handler"] == "subjects.math_10"
    assert sha256_file(registry.file_path(current)) == original


def test_rejected_candidate_is_not_activated(tmp_path):
    app = create_app(tmp_path)
    client = app.test_client()
    bad = tmp_path / "bad.docx"
    document = Document()
    document.add_table(rows=1, cols=1)
    document.save(str(bad))
    response = client.post(
        "/reference",
        data={"grade": "10", "subject": "Math", "reference": (bad.open("rb"), "bad.docx")},
        content_type="multipart/form-data",
    )
    assert "Reference Certification Failed." in response.get_data(as_text=True)
    assert ReferenceRegistry(tmp_path).rows() == []


def test_formatting_the_reference_keeps_its_text(tmp_path):
    output = tmp_path / "formatted.docx"
    result = format_exam(REFERENCE, REFERENCE, output)
    assert result.ok, result.message
    assert academic_text(read_exam(REFERENCE)) == academic_text(read_exam(output))
    assert resolved_labels(read_exam(REFERENCE)) == resolved_labels(read_exam(output))


def _inject(source: Path, dest: Path, snippet: str) -> None:
    shutil.copyfile(source, dest)
    with zipfile.ZipFile(dest) as package:
        root = etree.fromstring(package.read("word/document.xml"))
        payload = {name: package.read(name) for name in package.namelist()}
    paragraph = root.find(f".//{{{W}}}p")
    paragraph.append(etree.fromstring(snippet))
    payload["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(dest, "w") as package:
        for name, data in payload.items():
            package.writestr(name, data)

"""English 10 reports a missing section without renumbering the exam."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from subjects.english_10.service import format_exam, validate_exam

_HEADING = re.compile(r"^section\s*[-–—]?\s*([A-D])\s*$", re.IGNORECASE)

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _heading(text: str) -> str | None:
    match = _HEADING.match(" ".join(text.split()))
    if not match:
        return None
    return match.group(1).upper()
REFERENCE = FIXTURES / "good_reference.docx"
BAD = FIXTURES / "bad_teacher.docx"


def test_bad_exam_returns_yellow_error_fix_without_moving_section_d(tmp_path, monkeypatch):
    def explode(*_args, **_kwargs):
        raise AssertionError("hierarchy discovery ran for a nonconforming exam")

    monkeypatch.setattr("subjects.english_10.processor.discover._call_model", explode)
    output = tmp_path / "bad CORRECTIONS.docx"
    result = format_exam(REFERENCE, BAD, output)
    assert result.ok
    assert result.kind == "corrections"
    document = Document(str(output))
    texts = [paragraph.text.strip() for paragraph in document.paragraphs]
    errors = [text for text in texts if text.startswith("ERROR:")]
    fixes = [text for text in texts if text.startswith("FIX:")]
    assert any(text == "ERROR: Class header is missing." for text in errors)
    assert any(text == "ERROR: Section C is missing." for text in errors)
    assert any(text.startswith("FIX: Add the Class header") for text in fixes)
    assert any("Add Section C in the correct position" in text for text in fixes)
    headings = [_heading(text) for text in texts]
    assert "C" not in headings
    assert "D" in headings
    section_d = next(index for index, text in enumerate(texts) if _heading(text) == "D")
    section_c_error = next(index for index, text in enumerate(texts) if text == "ERROR: Section C is missing.")
    assert section_c_error < section_d
    assert "Ashoka" in "\n".join(texts)
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("ERROR:") or paragraph.text.strip().startswith("FIX:"):
            runs = [run for run in paragraph.runs if run.text]
            assert runs
            assert all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in runs)
            assert all(run.bold for run in runs)


def test_validate_reports_missing_section_without_failing_closed(monkeypatch):
    monkeypatch.setattr(
        "subjects.english_10.processor.discover._call_model",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("model called")),
    )
    result = validate_exam(REFERENCE, BAD)
    assert result.ok is False
    assert result.kind == "nonconformance"
    joined = "\n".join(result.diagnostics)
    assert "Section C is missing." in joined
    assert "Class header is missing." in joined

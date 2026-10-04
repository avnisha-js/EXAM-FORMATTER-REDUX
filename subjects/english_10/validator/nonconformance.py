"""Reference-exam conformance notes for a teacher Word file.

This does not format the exam and does not invent academic text.
A conforming exam returns no notes. A non-conforming exam is copied
and given ordinary yellow paragraphs the teacher can read and delete.
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from subjects.english_10.reference.headers import HeaderRequirement, load_header_policy
from subjects.english_10.reference.sections import canonical_heading, heading_key, required_section_headings


@dataclass(frozen=True)
class ExamError:
    code: str
    error: str
    fix: str
    place: str


def _norm(text: str) -> str:
    return " ".join((text or "").split())


def _present_headings(doc: Document, headings: list[str]) -> list[str]:
    found = []
    for paragraph in doc.paragraphs:
        canonical = canonical_heading(paragraph.text, headings)
        if canonical and canonical not in found:
            found.append(canonical)
    return found


def _before_first_section(doc: Document, headings: list[str]) -> list[str]:
    texts = []
    for paragraph in doc.paragraphs:
        if canonical_heading(paragraph.text, headings):
            break
        texts.append(paragraph.text)
    return texts


def _has_role(texts: list[str], requirement: HeaderRequirement) -> bool:
    return any(requirement.matches(text) for text in texts)


def _place_for_missing_section(present: list[str], heading: str, headings: list[str]) -> str:
    for later in headings[headings.index(heading) + 1 :]:
        if later in present:
            return f"before-section:{later}"
    return "bottom"


def find_nonconformances(teacher: Path, reference: Path) -> list[ExamError]:
    """Header roles and profile section headings the teacher lacks."""
    headings = required_section_headings(reference)
    requirements = load_header_policy(reference)
    teacher_doc = Document(str(teacher))
    teacher_header = _before_first_section(teacher_doc, headings)
    errors: list[ExamError] = []
    for requirement in requirements:
        if _has_role(teacher_header, requirement):
            continue
        errors.append(
            ExamError(
                f"missing_header_{requirement.role}",
                f"{requirement.label} header is missing.",
                f"Add the {requirement.label} header in the header area, following the Reference Exam.",
                "header",
            )
        )
    teacher_sections = _present_headings(teacher_doc, headings)
    seen_sections: list[str] = []
    for heading in headings:
        if heading in seen_sections:
            continue
        seen_sections.append(heading)
        if heading in teacher_sections:
            continue
        code = "missing_section_" + re.sub(r"\W+", "_", heading).strip("_")
        errors.append(
            ExamError(
                code,
                f"{heading} is missing.",
                f"Add {heading} in the correct position, following the Reference Exam.",
                _place_for_missing_section(teacher_sections, heading, headings),
            )
        )
    return errors


def _paint(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.bold = True
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _insert_before(anchor, text: str):
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    _paint(paragraph, text)
    return paragraph


def _append(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    _paint(paragraph, text)
    return paragraph


def _lines(errors: list[ExamError]) -> list[str]:
    lines = []
    for item in errors:
        lines.append(f"ERROR: {item.error}")
        lines.append(f"FIX: {item.fix}")
    return lines


def _insert_lines_before(anchor, lines: list[str]) -> None:
    # addprevious always inserts immediately before the anchor, so the
    # first line stays above the lines inserted after it.
    for line in lines:
        _insert_before(anchor, line)


def _first_header_paragraph(doc: Document, requirements: list[HeaderRequirement]):
    for paragraph in doc.paragraphs:
        if any(requirement.matches(paragraph.text) for requirement in requirements):
            return paragraph
    return doc.paragraphs[0] if doc.paragraphs else None


def _section_anchors(doc: Document) -> dict[tuple[str, ...], object]:
    anchors = {}
    for paragraph in doc.paragraphs:
        key = heading_key(paragraph.text)
        if key and key not in anchors:
            anchors[key] = paragraph
    return anchors


def write_nonconformance_file(teacher: Path, output: Path, errors: list[ExamError], reference: Path) -> Path:
    """Copy the teacher exam and insert yellow ERROR/FIX paragraphs."""
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(teacher, output)
    doc = Document(str(output))
    header = [item for item in errors if item.place == "header"]
    bottom = [item for item in errors if item.place == "bottom"]
    if header:
        anchor = _first_header_paragraph(doc, load_header_policy(reference))
        if anchor is None:
            bottom = header + bottom
        else:
            _insert_lines_before(anchor, _lines(header))
    anchors = _section_anchors(doc)
    for item in errors:
        if not item.place.startswith("before-section:"):
            continue
        heading = item.place.split(":", 1)[1]
        anchor = anchors.get(heading_key(heading))
        if anchor is None:
            bottom.append(item)
        else:
            _insert_lines_before(anchor, _lines([item]))
    for item in bottom:
        for line in _lines([item]):
            _append(doc, line)
    doc.save(str(output))
    return output

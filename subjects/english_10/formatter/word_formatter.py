"""Render the canonical hierarchy as a printable Word exam.

The hierarchy is consumed as-is. Teacher wording is kept. Only structural
labels, spacing, and copying each teacher image at its extracted host
are presentation.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_UNDERLINE
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Length, Pt, Twips

from subjects.english_10.formatter.styles import (
    INDENT_SYLLABUS,
    MAX_IMAGE_WIDTH,
    SPACE_OR_AFTER,
    SPACE_OR_BEFORE,
    SPACE_PASSAGE_AFTER,
    SPACE_PASSAGE_BEFORE,
    SPACE_SYLLABUS_HEAD_AFTER,
    SPACE_SYLLABUS_HEAD_BEFORE,
)
from subjects.english_10.reference.appearance import Appearance, RoleStyle, load_appearance
from subjects.english_10.reference.headers import load_header_policy
from subjects.english_10.reference.numbering import load_numbering_policy
from subjects.english_10.reference.sections import canonical_heading, required_section_headings
from subjects.english_10.reference.syllabus import load_syllabus_policy

_RULE = re.compile(r"^[\s_\-–—]+$")
_CHOICE_MARK = re.compile(r"([A-Da-d])\s*[\.\)]")


def strip_leading_marker(text: str, policy) -> str:
    """Remove one leading label from the active numbering policy."""
    return policy.strip_leading(text)


def strip_trailing_marker(text: str, marker: str | None) -> str:
    """Remove one trailing alternative marker when the profile defines one."""
    if not marker:
        return (text or "").rstrip()
    pattern = re.compile(rf"\s*\b{re.escape(marker)}\s*$")
    return pattern.sub("", text or "", count=1).rstrip()


def choice_pieces(text: str, policy) -> list[str]:
    """Split one choice block into option texts, in marker order."""
    if not policy.choice:
        return [(text or "").strip()]
    if policy.source != "compatibility":
        return [(text or "").strip()]
    found: list[tuple[str, int, int]] = []
    raw = text or ""
    for match in _CHOICE_MARK.finditer(raw):
        letter = match.group(1).lower()
        prev = raw[match.start() - 1] if match.start() else ""
        if not found:
            if letter not in "abcd":
                continue
        elif letter != chr(ord(found[-1][0]) + 1):
            continue
        elif prev.isalpha() and letter == "a":
            continue
        found.append((letter, match.start(), match.end()))
        if letter == "d":
            break
    if not found:
        return [raw.strip()]
    pieces = []
    for index, (_letter, _start, end) in enumerate(found):
        stop = found[index + 1][1] if index + 1 < len(found) else len(raw)
        pieces.append(raw[end:stop].strip())
    return pieces


def _ordered(blocks):
    return sorted(
        blocks,
        key=lambda b: (b.source_order, 0 if b.source_id.startswith("p") else 1, b.source_id),
    )


def _is_rule(text: str) -> bool:
    return bool(_RULE.fullmatch(text or ""))


def _media_name(media_part: str) -> str:
    part = media_part.replace("\\", "/")
    if part.startswith("word/"):
        return part
    return "word/" + part.lstrip("/")


def _image_bytes_and_width(teacher_path: Path, media_part: str):
    name = _media_name(media_part)
    with zipfile.ZipFile(teacher_path) as zf:
        data = zf.read(name)
        root_xml = zf.read("word/document.xml")
        width = _default_width()
    from lxml import etree

    root = etree.fromstring(root_xml)
    ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    extent = root.find(".//wp:extent", ns)
    if extent is not None and extent.get("cx"):
        width = Emu(int(extent.get("cx")))
    if width > MAX_IMAGE_WIDTH:
        width = MAX_IMAGE_WIDTH
    return data, width


def _default_width():
    return Inches(3.0)


def _apply_run(paragraph, text: str, size_pt, font_name, bold: bool = False, underline: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.underline = WD_UNDERLINE.SINGLE if underline else WD_UNDERLINE.NONE
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if font_name:
        run.font.name = font_name
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        fonts.set(qn("w:ascii"), font_name)
        fonts.set(qn("w:hAnsi"), font_name)
    return run


def _paragraph(
    doc,
    text,
    appearance: Appearance,
    role: RoleStyle | None,
    *,
    before=None,
    after=None,
    left=None,
    align=None,
    bold=None,
    underline=False,
    size_pt=None,
    font_name=None,
):
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    space_before = before if before is not None else (role.space_before_pt if role and role.space_before_pt is not None else 0)
    space_after = after if after is not None else (role.space_after_pt if role and role.space_after_pt is not None else 0)
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    if appearance.line_spacing_multiple is not None:
        fmt.line_spacing = appearance.line_spacing_multiple
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    indent = left if left is not None else (role.left_indent_in if role else None)
    if isinstance(indent, Length):
        if indent:
            fmt.left_indent = indent
    elif isinstance(indent, (int, float)) and indent:
        fmt.left_indent = Inches(indent)
    if align is not None:
        paragraph.alignment = align
    if text:
        use_bold = bold if bold is not None else (bool(role.bold) if role and role.bold is not None else False)
        use_font = font_name if font_name is not None else appearance.font_for(role)
        use_size = size_pt if size_pt is not None else appearance.size_for(role)
        _apply_run(paragraph, text, use_size, use_font, bold=use_bold, underline=underline)
    return paragraph


def _prepare_document(appearance: Appearance) -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    if appearance.predominant_font:
        normal.font.name = appearance.predominant_font
    if appearance.default_size_pt is not None:
        normal.font.size = Pt(appearance.default_size_pt)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    if appearance.line_spacing_multiple is not None:
        normal.paragraph_format.line_spacing = appearance.line_spacing_multiple
    section = doc.sections[0]
    if appearance.orientation == "LANDSCAPE":
        section.orientation = WD_ORIENT.LANDSCAPE
    elif appearance.orientation == "PORTRAIT":
        section.orientation = WD_ORIENT.PORTRAIT
    if appearance.page_width_twip is not None:
        section.page_width = Twips(appearance.page_width_twip)
    if appearance.page_height_twip is not None:
        section.page_height = Twips(appearance.page_height_twip)
    if appearance.margin_left_twip is not None:
        section.left_margin = Twips(appearance.margin_left_twip)
    if appearance.margin_right_twip is not None:
        section.right_margin = Twips(appearance.margin_right_twip)
    if appearance.margin_top_twip is not None:
        section.top_margin = Twips(appearance.margin_top_twip)
    if appearance.margin_bottom_twip is not None:
        section.bottom_margin = Twips(appearance.margin_bottom_twip)
    return doc


def _indent(block, by_id, implicit_major_id, appearance: Appearance):
    extra = 0
    parent = by_id.get(block.parent_id)
    seen = set()
    while parent is not None and parent.source_id not in seen:
        seen.add(parent.source_id)
        if parent.block_type == "branch":
            extra += 1
        parent = by_id.get(parent.parent_id)
    if (
        implicit_major_id
        and block.parent_id == implicit_major_id
        and block.block_type != "branch"
    ):
        extra += 1
    step = appearance.role("subquestion").left_indent_in or 0
    if block.block_type == "choice":
        base = appearance.role("choice").left_indent_in or 0
    elif block.block_type in {"passage", "subquestion", "continuation", "alternative", "branch"}:
        base = step
    else:
        base = 0
    return base + (step * extra)


def _header_lines(blocks, requirements, syllabus_heading: str | None = None):
    """Metadata run that contains a header role required by the profile.

    The raw teacher file may store that cluster after the syllabus list. A
    formatted file stores it above the syllabus. Keep the cluster in either
    place. Leave other pre-syllabus metadata out.
    """
    ordered = _ordered(blocks)
    runs: list[list[str]] = []
    current: list[str] = []

    def close() -> None:
        nonlocal current
        if current:
            runs.append(current)
            current = []

    for block in ordered:
        text = " ".join((block.original_text or "").split())
        if block.block_type == "metadata" and text and not _is_rule(block.original_text):
            if syllabus_heading and text.casefold() == syllabus_heading.casefold():
                close()
                continue
            current.append(block.original_text.strip())
            continue
        if text and block.block_type != "metadata":
            close()
    close()
    anchored = [run for run in runs if any(requirement.matches(line) for requirement in requirements for line in run)]
    if anchored:
        return anchored[0]
    return []


def _section_label(text: str, headings: list[str]) -> str:
    canonical = canonical_heading(text, headings)
    if canonical:
        return canonical
    return (text or "").strip()


def _header_line_style(appearance: Appearance, index: int, count: int) -> tuple[RoleStyle, float | None]:
    lines = appearance.header_lines
    if not lines:
        return RoleStyle(), None
    source = lines[index] if index < len(lines) else lines[-1]
    after = source.space_after_pt
    if index == count - 1 and lines[-1].space_after_pt is not None:
        after = lines[-1].space_after_pt
    return source, after


def format_exam(blocks, teacher_path: Path, out_path: Path, reference_path: Path) -> Path:
    section_headings = required_section_headings(reference_path)
    policy = load_numbering_policy(reference_path)
    header_requirements = load_header_policy(reference_path)
    appearance = load_appearance(reference_path)
    syllabus_policy = load_syllabus_policy(reference_path)
    by_id = {b.source_id: b for b in blocks}
    doc = _prepare_document(appearance)
    header = _header_lines(blocks, header_requirements, syllabus_policy.heading)
    printed_header = {" ".join(line.split()) for line in header}
    major = appearance.role("major_question")
    section_role = appearance.role("section_heading")
    sub_role = appearance.role("subquestion")
    choice_role = appearance.role("choice")
    body = appearance.role("body")
    for index, line in enumerate(header):
        source, after = _header_line_style(appearance, index, len(header))
        _paragraph(
            doc,
            line,
            appearance,
            source,
            before=0,
            after=after,
            bold=True,
            font_name=appearance.font_for(source),
            size_pt=appearance.size_for(source),
        )

    step = sub_role.left_indent_in or 0
    syllabus = [b for b in _ordered(blocks) if b.block_type == "syllabus"]
    if syllabus and syllabus_policy.heading:
        _paragraph(
            doc,
            syllabus_policy.heading,
            appearance,
            major,
            before=SPACE_SYLLABUS_HEAD_BEFORE.pt,
            after=SPACE_SYLLABUS_HEAD_AFTER.pt,
            bold=True,
        )
    if syllabus:
        for block in syllabus:
            _paragraph(
                doc,
                block.original_text.strip(),
                appearance,
                body,
                before=0,
                after=1,
                left=INDENT_SYLLABUS,
            )

    implicit_major_id = None
    seen_section = False
    for block in _ordered(blocks):
        kind = block.block_type
        if kind in {"metadata", "other"} and not seen_section:
            text = " ".join((block.original_text or "").split())
            if not text or _is_rule(block.original_text):
                continue
            if text in printed_header:
                continue
            if syllabus_policy.heading and text.casefold() == syllabus_policy.heading.casefold():
                continue
            _paragraph(doc, block.original_text.strip(), appearance, body)
            continue
        if kind in {"metadata", "syllabus", "other"}:
            continue
        if kind == "section":
            seen_section = True
            implicit_major_id = None
            _paragraph(
                doc,
                _section_label(block.original_text, section_headings),
                appearance,
                section_role,
                bold=True,
                underline=True,
            )
            continue
        if kind == "major_question":
            implicit_major_id = block.source_id if block.branch_carrier else None
            stem = strip_trailing_marker(strip_leading_marker(block.original_text, policy), policy.alternative)
            label = block.final_label or ""
            q_label = label.split()[0] if label and label != "-" else ""
            paragraph = _paragraph(doc, "", appearance, major)
            major_font = appearance.font_for(major)
            major_size = appearance.size_for(major)
            if q_label:
                _apply_run(paragraph, q_label, major_size, major_font, bold=True, underline=True)
                if stem:
                    _apply_run(paragraph, " " + stem, major_size, major_font, bold=False, underline=False)
            elif stem:
                _apply_run(paragraph, stem, major_size, major_font, bold=False, underline=False)
            if block.branch_carrier and policy.branch:
                branch_label = label.split()[1] if len(label.split()) > 1 else policy.label("branch", 1)
                _paragraph(
                    doc,
                    branch_label,
                    appearance,
                    major,
                    before=2,
                    after=2,
                    left=step,
                    bold=True,
                )
            if policy.alternative and re.search(rf"\s*\b{re.escape(policy.alternative)}\s*$", block.original_text or ""):
                _paragraph(
                    doc,
                    policy.alternative,
                    appearance,
                    major,
                    before=SPACE_OR_BEFORE.pt,
                    after=SPACE_OR_AFTER.pt,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    bold=True,
                )
            continue
        if kind == "or_marker":
            if policy.alternative:
                _paragraph(
                    doc,
                    policy.alternative,
                    appearance,
                    major,
                    before=SPACE_OR_BEFORE.pt,
                    after=SPACE_OR_AFTER.pt,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    bold=True,
                )
            elif (block.original_text or "").strip():
                _paragraph(doc, block.original_text.strip(), appearance, body)
            continue
        if kind == "image":
            data, width = _image_bytes_and_width(teacher_path, block.media_part)
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.add_run().add_picture(BytesIO(data), width=width)
            continue
        if kind == "branch":
            passage = strip_leading_marker(block.original_text, policy)
            shown = block.final_label if block.final_label and block.final_label != "-" else ""
            if not shown and policy.branch and policy.source == "compatibility":
                shown = policy.label("branch", 2) or ""
            _paragraph(
                doc,
                shown,
                appearance,
                major,
                after=2,
                left=step,
                bold=True,
            )
            if passage:
                _paragraph(
                    doc,
                    passage,
                    appearance,
                    body,
                    before=SPACE_PASSAGE_BEFORE.pt,
                    after=SPACE_PASSAGE_AFTER.pt,
                    left=step + step,
                )
            continue
        left = _indent(block, by_id, implicit_major_id, appearance)
        if kind == "choice":
            labels = (block.final_label or "").split()
            pieces = choice_pieces(block.original_text, policy)
            if len(labels) != len(pieces):
                fallback = policy.label("choice", 1) if policy.choice else ""
                labels = labels or ([fallback] if fallback else [])
                while len(labels) < len(pieces):
                    labels.append(labels[-1])
            choice_style = choice_role if choice_role.font_size_pt is not None else body
            for label, piece in zip(labels, pieces):
                _paragraph(
                    doc,
                    f"{label} {piece}".strip(),
                    appearance,
                    choice_style,
                    left=left,
                )
            continue
        if kind == "subquestion":
            sub_body = strip_leading_marker(block.original_text, policy)
            label = block.final_label if block.final_label and block.final_label != "-" else ""
            shown = f"{label} {sub_body}".strip()
            _paragraph(doc, shown, appearance, sub_role, left=left)
            continue
        if kind in {"passage", "continuation", "alternative"}:
            if kind == "passage":
                before, after = SPACE_PASSAGE_BEFORE.pt, SPACE_PASSAGE_AFTER.pt
            else:
                before = sub_role.space_before_pt if sub_role.space_before_pt is not None else 0
                after = sub_role.space_after_pt if sub_role.space_after_pt is not None else 0
            _paragraph(
                doc,
                block.original_text.strip(),
                appearance,
                body,
                before=before,
                after=after,
                left=left,
            )
            continue
        _paragraph(doc, block.original_text.strip(), appearance, body, left=left)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path

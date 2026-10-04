"""Check the formatted DOCX against the canonical hierarchy and teacher text."""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document

from subjects.english_10.formatter.word_formatter import choice_pieces, strip_leading_marker, strip_trailing_marker
from subjects.english_10.reference.images import load_image_policy
from subjects.english_10.reference.numbering import load_numbering_policy
from subjects.english_10.reference.sections import contains_heading, load_profile_for_reference, required_section_headings

_MARK = re.compile(r"\(\s*\d+\s*x\s*\d+\s*=\s*\d+\s*\)|\(\s*\d+\s*\)")
_WS = re.compile(r"\s+")


def _norm(text: str) -> str:
    return _WS.sub(" ", text or "").strip()


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _paragraphs(path: Path) -> list[str]:
    doc = Document(str(path))
    return [p.text for p in doc.paragraphs]


def _image_hashes(path: Path) -> list[str]:
    hashes = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/"):
                hashes.append(hashlib.sha256(zf.read(name)).hexdigest())
    return hashes


def _body(block, policy) -> str:
    kind = block.block_type
    text = block.original_text or ""
    if kind == "major_question":
        return strip_trailing_marker(strip_leading_marker(text, policy), policy.alternative)
    if kind in {"subquestion", "branch"}:
        return strip_leading_marker(text, policy)
    if kind == "choice":
        return text
    return text.strip()


def check_format(blocks, teacher_path: Path, formatted_path: Path, reference_path: Path) -> list[str]:
    problems: list[str] = []
    if not formatted_path.exists():
        return ["Formatted DOCX was not created."]
    raw = _docx_text(formatted_path)
    flat = _norm(raw)
    paragraphs = _paragraphs(formatted_path)
    policy = load_numbering_policy(reference_path)

    majors = sorted(
        (b for b in blocks if b.block_type == "major_question"),
        key=lambda b: b.source_order,
    )
    cursor = 0
    for index, major in enumerate(majors, start=1):
        token = policy.label("major", index) or ""
        found = flat.find(token, cursor)
        if found < 0:
            problems.append(f"Missing {token}")
        else:
            cursor = found + len(token)
        body = _norm(_body(major, policy))
        if body and body not in flat:
            problems.append(f"{token} wording missing: {body[:80]}")

    for heading in required_section_headings(reference_path):
        if not contains_heading(flat, heading):
            problems.append(f"Missing {heading}")

    for block in blocks:
        kind = block.block_type
        if kind == "subquestion":
            label = block.final_label
            body = _norm(_body(block, policy))
            shown = _norm(f"{label} {body}")
            if shown not in flat:
                problems.append(f"Subquestion missing {block.source_id}: {shown[:80]}")
        elif kind == "passage":
            body = _norm(block.original_text)
            if body and body not in flat:
                problems.append(f"Passage missing {block.source_id}")
        elif kind == "alternative":
            body = _norm(block.original_text)
            if body and body not in flat:
                problems.append(f"Alternative missing {block.source_id}")
        elif kind == "choice":
            for piece in choice_pieces(block.original_text, policy):
                if piece and _norm(piece) not in flat:
                    problems.append(f"Choice text missing {block.source_id}: {piece[:60]}")
        elif kind == "continuation":
            body = _norm(block.original_text)
            if body and body not in flat:
                problems.append(f"Continuation missing {block.source_id}")
        elif kind == "branch":
            body = _norm(_body(block, policy))
            label = block.final_label or ""
            if not label and policy.branch and policy.source == "compatibility":
                label = policy.label("branch", 2) or ""
            if label not in flat:
                problems.append(f"Branch label missing {block.source_id}")
            if body and body not in flat:
                problems.append(f"Branch passage missing {block.source_id}")
        elif kind == "syllabus":
            body = _norm(block.original_text)
            if body and body not in flat:
                problems.append(f"Syllabus line missing {block.source_id}")
            for paragraph in paragraphs:
                numbered = paragraph.strip().startswith("Q") if policy.source == "compatibility" else False
                if _norm(paragraph) == body and numbered:
                    problems.append(f"Syllabus line numbered {block.source_id}")

    source_text = "\n".join(b.original_text or "" for b in blocks)
    for mark in _MARK.findall(source_text):
        if _norm(mark) not in flat:
            problems.append(f"Mark missing: {mark}")

    for block in blocks:
        if block.block_type in {"metadata", "syllabus", "passage", "alternative", "continuation"}:
            for number in re.findall(r"\d+", block.original_text or ""):
                if number not in flat:
                    problems.append(f"Number {number} missing from {block.source_id}")
                    break
        elif block.block_type in {"major_question", "subquestion", "branch", "choice"}:
            for number in re.findall(r"\d+", _body(block, policy)):
                if number not in flat:
                    problems.append(f"Number {number} missing from {block.source_id}")
                    break

    teacher_hashes = _image_hashes(teacher_path)
    out_hashes = _image_hashes(formatted_path)
    image_policy = load_image_policy(reference_path)
    if len(out_hashes) != image_policy.count:
        problems.append(f"Expected {image_policy.count} image(s), found {len(out_hashes)}")
    elif image_policy.count and teacher_hashes and any(digest not in teacher_hashes for digest in out_hashes):
        problems.append("Formatted image is not the teacher exam image")

    teacher_flat = _norm(_docx_text(teacher_path))
    profile = load_profile_for_reference(reference_path)
    header = profile.get("header") or {}
    header_lines = set()
    if header.get("status") == "observed":
        for line in header.get("lines") or []:
            text = _norm(line.get("text") or "")
            if text:
                header_lines.add(text)
    reference = Document(str(reference_path))
    for paragraph in reference.paragraphs:
        snippet = _norm(paragraph.text)
        if not snippet or snippet not in flat:
            continue
        if snippet in teacher_flat:
            continue
        if len(snippet) >= 60:
            problems.append(f"Reference wording leaked: {snippet[:80]}")
        elif snippet in header_lines:
            problems.append(f"Reference wording leaked: {snippet[:80]}")
    return problems
